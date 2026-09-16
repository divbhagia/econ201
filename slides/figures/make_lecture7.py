"""
Figures and numbers for Lecture 7: the class's demand for a day at a Dodgers
game, built from their own answers, and the elasticity of demand along it
(CORE 7.5), plus the two CORE question figures on the practice page.

Run from the course root:

    python3 slides/figures/make_lecture7.py

The class demand curve comes from the Lecture 3 Mentimeter poll, saved as
slides/figures/data/menti-lec3.xlsx (anonymous voter numbers only). It uses
the Dodgers question from the in-class session, session 2 on 2026-08-31;
session 1 is a single test answer and is left out. Blank answers are dropped.

The straight line is drawn by eye with round numbers, not fitted: least
squares on the 88 answers gives Q = 79 - 0.25 P, pulled around by the few
$400 answers, while Q = 90 - 0.3 P (P = 300 - (10/3) Q) stays within 3
students of the actual counts at $50, $100, $150, and $250. Every $10 rise
in price costs 3 students, and the elasticity -(P/Q)(dQ/dP) = 0.3 P / Q is
0.2 at $50, 0.5 at $100, and 2 at $200.

Output, in slides/img/:
  class-demand.svg   the class's answers lined up highest first, as a
                     demand curve, $100 and $200 marked
  class-linear.svg   the same answers in grey with the straight line drawn
                     through them, labelled in its P = form, $100 marked
  class-elasticity.svg  the straight line alone, the three Activity prices
                     ($50, $100, $200) marked with their elasticities
  flat-steep.svg     a flatter and a steeper demand curve through the same
                     point, $100 and 60 students, with their elasticities
  core-q7-6.svg      CORE Question 7.6, demand curves D and D'
  core-q7-8.svg      CORE Question 7.8, demand curves D1 and D2

Accessibility: Okabe-Ito derived colours, every marked point labelled with
its value on the figure itself, so nothing rests on reading positions or
colours by eye; text converted to paths so the SVG looks the same on any
machine. Alt text lives in the slide and practice files beside each image.
"""
import argparse
import glob as _glob
import os
import sys
from pathlib import Path

import matplotlib
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.ticker import FuncFormatter
from openpyxl import load_workbook

matplotlib.use("Agg")

# Same font setup as make_lecture6.py.
for f in Path.home().glob("Library/Fonts/FiraSans-*.otf"):
    font_manager.fontManager.addfont(str(f))
for pat in ("/usr/local/texlive/*/texmf-dist/fonts/truetype/typoland/lato/*.ttf",
            str(Path.home() / "Library/Fonts/Lato-*.ttf")):
    for f in _glob.glob(pat):
        font_manager.fontManager.addfont(f)

ROOT = Path(__file__).resolve().parents[2]
IMG = ROOT / "slides" / "img"
MENTI = ROOT / "slides" / "figures" / "data" / "menti-lec3.xlsx"
MENTI_SESSION = "2"
CLASS_PRICES = (100, 200)    # marked on the class demand curve
# The straight-line version, Q = LIN_A - LIN_B P.
LIN_A, LIN_B = 90, 0.3
LIN_STEP = 10                # the $10 price rises in Activity 1
LIN_EXAMPLE = 100            # worked on the slides
LIN_ACTIVITY = (50, 200)     # Activity 1: percentage changes
POINT_ACTIVITY = (50, 100, 200)  # Activity 2: -(P/Q)(dQ/dP), same prices as Activity 1
# Flat versus steep: two lines through ($100, 60 students), Q = 60 + b (100 - P),
# so the elasticity at $100 is b * 100 / 60.
FLAT_B, STEEP_B = 0.9, 0.15
SCRATCH = Path(os.environ["FIG_PREVIEW_DIR"]) if os.environ.get("FIG_PREVIEW_DIR") else None

INK = "#1a1a1a"
GRID = "#d9d9d9"
DEMAND = "#0072B2"
DEMAND2 = "#BF5700"
GUIDE = "#595a5b"
CLASS = "#BF5700"            # brand orange, 4.59:1 on white
CLASS_SIZE = (10.8, 4.8)
LINEAR_SIZE = (8.6, 5.6)     # taller and narrower, for the straight-line slide
SIDE = (8.2, 5.9)
FONT = 22

plt.rcParams.update(
    {
        "font.family": ["Lato", "Fira Sans", "sans-serif"],
        "font.size": FONT,
        "svg.fonttype": "path",
        "pdf.fonttype": 42,
        "axes.edgecolor": INK,
        "axes.labelcolor": INK,
        "xtick.color": INK,
        "ytick.color": INK,
        "xtick.labelsize": FONT,
        "ytick.labelsize": FONT,
    }
)

def save(fig, out: Path) -> None:
    IMG.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, format="svg")
    if SCRATCH:
        fig.savefig(SCRATCH / f"{out.stem}.png", format="png", dpi=110)
    plt.close(fig)
    print(f"wrote {out.relative_to(ROOT)}")


def guides(ax, x, y):
    ax.plot([x, x], [0, y], color=GUIDE, lw=1.6, ls=(0, (4, 3)), zorder=4)
    ax.plot([0, x], [y, y], color=GUIDE, lw=1.6, ls=(0, (4, 3)), zorder=4)


def mark_xy(ax, x, y, text, dx=14, dy=10, ha="left", va="bottom", guide=False,
            color=DEMAND):
    if guide:
        guides(ax, x, y)
    ax.plot(x, y, marker="o", ms=13, color=color, zorder=7)
    ax.annotate(text, xy=(x, y), xytext=(dx, dy),
                textcoords="offset points", ha=ha, va=va,
                fontsize=FONT, color=color, fontweight="bold", linespacing=1.2, zorder=8)


def class_wtp():
    """Willingness to pay for the Dodgers day, highest first, from the poll."""
    rows = list(load_workbook(MENTI, read_only=True)["Voters"].iter_rows(values_only=True))
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Date (UTC)")
    col = next(j for j, h in enumerate(rows[header]) if h and "Dodgers" in str(h))
    return sorted((float(r[col]) for r in rows[header + 1:]
                   if str(r[1]) == MENTI_SESSION and r[col] is not None), reverse=True)


def buyers(wtp, p):
    return sum(w >= p for w in wtp)


def draw_class_demand(out: Path) -> None:
    """Each answer is one step: at any price, demand is everyone at or above it.

    Orange rather than the demand blue, at Div's request, to set the class's
    own answers apart from the Beautiful Cars curves. Wider than DEMAND_SIZE,
    so the figure fills more of the slide under its two lines of text.
    """
    wtp = class_wtp()
    n = len(wtp)
    fig, ax = plain_axes(CLASS_SIZE, 95, 450)
    fig.subplots_adjust(left=0.115, right=0.975, top=0.95, bottom=0.18)
    ax.set_xticks(range(0, 91, 10))
    ax.set_yticks(range(0, 401, 100))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"${y:,.0f}"))
    x = np.arange(0, n + 1)
    ax.step(x, wtp + [wtp[-1]], where="post", color=CLASS, lw=2.8, zorder=5)
    ax.annotate("Demand", xy=(5, 400), xytext=(12, 0), textcoords="offset points",
                ha="left", va="center", fontsize=FONT, color=CLASS,
                fontweight="bold", zorder=8)
    for p in CLASS_PRICES:
        k = buyers(wtp, p)
        mark_xy(ax, k, p, f"{k} students at ${p}", guide=True, color=CLASS)
    ax.set_xlabel("Number of students", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price for the day", fontsize=FONT, labelpad=8)
    save(fig, out)


def class_linear(p):
    return LIN_A - LIN_B * p


def draw_class_linear(out: Path) -> None:
    """The class's answers in grey, with the straight line through them."""
    wtp = class_wtp()
    n = len(wtp)
    fig, ax = plain_axes(LINEAR_SIZE, 95, 450)
    fig.subplots_adjust(left=0.145, right=0.965, top=0.95, bottom=0.15)
    ax.set_xticks(range(0, 91, 10))
    ax.set_yticks(range(0, 401, 100))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"${y:,.0f}"))
    x = np.arange(0, n + 1)
    ax.step(x, wtp + [wtp[-1]], where="post", color="#9a9a9a", lw=2.0, zorder=4)
    ax.annotate("Your answers", xy=(5, 400), xytext=(12, 0), textcoords="offset points",
                ha="left", va="center", fontsize=FONT, color=GUIDE, zorder=8)
    pr = np.array([0, LIN_A / LIN_B])
    ax.plot(class_linear(pr), pr, color=CLASS, lw=3.2, solid_capstyle="round", zorder=5)
    # Labelled in the P = form, to match the axes; the slide then solves for Q.
    ax.annotate(f"P = {LIN_A / LIN_B:g} − (10/3)Q", xy=(class_linear(260), 260), xytext=(14, 6),
                textcoords="offset points", ha="left", va="bottom", fontsize=FONT,
                color=CLASS, fontweight="bold", zorder=8)
    p0 = LIN_EXAMPLE
    mark_xy(ax, class_linear(p0), p0, f"{class_linear(p0):g} students\nat ${p0}",
            guide=True, color=CLASS)
    ax.set_xlabel("Number of students", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price for the day", fontsize=FONT, labelpad=8)
    save(fig, out)


def draw_class_elasticity(out: Path) -> None:
    """The straight line with the Activity prices marked by their elasticity.

    Prices sit on $50 gridlines with a dotted guide from the price axis, and
    every label is to the right of the line.
    """
    fig, ax = plain_axes(SIDE, 95, 320)
    fig.subplots_adjust(left=0.16, right=0.965, top=0.96, bottom=0.15)
    ax.set_xticks(range(0, 91, 15))
    ax.set_yticks(range(0, 301, 50))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"${y:,.0f}"))
    pr = np.array([0, LIN_A / LIN_B])
    ax.plot(class_linear(pr), pr, color=CLASS, lw=3.2, solid_capstyle="round", zorder=5)
    for p0 in POINT_ACTIVITY:
        q0 = class_linear(p0)
        eps = LIN_B * p0 / q0
        ax.plot([0, q0], [p0, p0], color=CLASS, lw=2.6, ls=(0, (5, 3)), zorder=4)
        ax.plot(q0, p0, marker="o", ms=13, color=CLASS, zorder=7)
        ax.annotate(f"ε = {eps:g}", xy=(q0, p0), xytext=(16, 6), textcoords="offset points",
                    ha="left", va="bottom", fontsize=FONT + 6, color=CLASS,
                    fontweight="bold", zorder=8)
    ax.set_xlabel("Number of students", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price for the day", fontsize=FONT, labelpad=8)
    save(fig, out)


def draw_flat_steep(out: Path) -> None:
    """Your demand curve and a flatter one through the same point.

    Both pass through 60 buyers at $100. The same rise to $150 cuts quantity
    to 45 on your curve and to 15 on the flatter one, so at $100 demand is
    more elastic on the flatter curve (1.5 against 0.5).
    """
    fig, ax = plain_axes(SIDE, 160, 320)
    fig.subplots_adjust(left=0.16, right=0.965, top=0.96, bottom=0.15)
    ax.set_xticks(range(0, 151, 15))
    ax.set_xticklabels([str(t) if t % 30 == 0 or t in (15, 45) else "" for t in range(0, 151, 15)])
    ax.set_yticks(range(0, 301, 50))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"${y:,.0f}"))
    q0, p0, p1 = 60, 100, 150
    curves = ((LIN_B, CLASS, "Your demand"), (FLAT_B, DEMAND, "Flatter demand"))
    for b, colour, name in curves:
        top = p0 + q0 / b
        prices = np.array([0, min(top, 320)])
        ax.plot(q0 + b * (p0 - prices), prices, color=colour, lw=3, zorder=5)
    # the price rise, with dotted guides to both axes
    ax.plot([0, q0], [p0, p0], color=GUIDE, lw=1.6, ls=(0, (4, 3)), zorder=4)
    ax.plot([q0, q0], [0, p0], color=GUIDE, lw=1.6, ls=(0, (4, 3)), zorder=4)
    ax.plot([0, q0 + LIN_B * (p0 - p1)], [p1, p1], color=GUIDE, lw=1.6, ls=(0, (4, 3)), zorder=4)
    ax.plot(q0, p0, marker="o", ms=13, color=INK, zorder=7)
    for b, colour, name in curves:
        q1 = q0 + b * (p0 - p1)
        ax.plot([q1, q1], [0, p1], color=colour, lw=1.4, ls=(0, (4, 3)), zorder=4)
        ax.plot(q1, p1, marker="o", ms=12, color=colour, zorder=7)
    # Curve names in open space; the elasticities belong to the $100 point.
    ax.text(8, 290, "Your demand", ha="left", va="center", fontsize=FONT,
            color=CLASS, fontweight="bold", zorder=8)
    ax.text(96, 72, "Flatter demand", ha="left", va="center", fontsize=FONT,
            color=DEMAND, fontweight="bold", zorder=8)
    ax.annotate(f"At $100:", xy=(q0, p0), xytext=(18, 58), textcoords="offset points",
                ha="left", va="bottom", fontsize=FONT - 2, color=INK, zorder=8)
    ax.annotate(f"ε = {LIN_B * p0 / q0:g} on yours", xy=(q0, p0), xytext=(18, 32),
                textcoords="offset points", ha="left", va="bottom", fontsize=FONT - 2,
                color=CLASS, fontweight="bold", zorder=8)
    ax.annotate(f"ε = {FLAT_B * p0 / q0:g} on the flatter one", xy=(q0, p0), xytext=(18, 6),
                textcoords="offset points", ha="left", va="bottom", fontsize=FONT - 2,
                color=DEMAND, fontweight="bold", zorder=8)
    ax.set_xlabel("Number of buyers", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price", fontsize=FONT, labelpad=8)
    save(fig, out)


def plain_axes(size, xmax, ymax):
    fig, ax = plt.subplots(figsize=size, dpi=100)
    ax.grid(color=GRID, lw=1, zorder=0)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(length=5, width=1)
    ax.set_xlim(0, xmax)
    ax.set_ylim(0, ymax)
    return fig, ax


def draw_q76(out: Path) -> None:
    """CORE Question 7.6: D from (0, 6,000) to (60, 0), D' from (0, 10,000) to (100, 0)."""
    fig, ax = plain_axes(SIDE, 105, 12_000)
    fig.subplots_adjust(left=0.22, right=0.965, top=0.93, bottom=0.15)
    ax.set_xticks(range(0, 101, 20))
    ax.set_yticks(range(0, 12_001, 2_000))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"${y:,.0f}"))
    ax.plot([0, 60], [6_000, 0], color=DEMAND, lw=2.6, zorder=5)
    ax.plot([0, 100], [10_000, 0], color=DEMAND2, lw=2.6, zorder=5)
    ax.annotate("D", xy=(40, 2_000), xytext=(-10, -10), textcoords="offset points",
                ha="right", va="top", fontsize=FONT, color=DEMAND, fontweight="bold")
    ax.annotate("D′", xy=(70, 3_000), xytext=(10, 10), textcoords="offset points",
                ha="left", va="bottom", fontsize=FONT, color=DEMAND2, fontweight="bold")
    ax.set_xlabel("Quantity", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price", fontsize=FONT, labelpad=8)
    save(fig, out)


def draw_q78(out: Path) -> None:
    """CORE Question 7.8: steep D1 and flat D2 crossing at E; A, C on D1, B on D2."""
    fig, ax = plain_axes(SIDE, 10, 10)
    fig.subplots_adjust(left=0.09, right=0.965, top=0.96, bottom=0.11)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.grid(False)
    # D1: P = 9 - 1.2 Q, D2: P = 6.5 - 0.35 Q, crossing at E.
    def d1(q):
        return 9 - 1.2 * q

    def d2(q):
        return 6.5 - 0.35 * q

    qe = 2.5 / 0.85
    ax.plot([0, 7.5], [d1(0), d1(7.5)], color=DEMAND, lw=2.6, zorder=5)
    ax.plot([0, 10], [d2(0), d2(10)], color=DEMAND2, lw=2.6, zorder=5)
    ax.annotate("D₁", xy=(6.6, d1(6.6)), xytext=(10, 0), textcoords="offset points",
                ha="left", va="center", fontsize=FONT, color=DEMAND, fontweight="bold")
    ax.annotate("D₂", xy=(9.3, d2(9.3)), xytext=(0, 10), textcoords="offset points",
                ha="center", va="bottom", fontsize=FONT, color=DEMAND2, fontweight="bold")
    points = {"A": (1.2, d1(1.2)), "E": (qe, d1(qe)), "C": (5.2, d1(5.2)),
              "B": (7.0, d2(7.0))}
    for name, (x, y) in points.items():
        ax.plot(x, y, marker="o", ms=12, color=INK, zorder=7)
        ax.annotate(name, xy=(x, y), xytext=(12, 8), textcoords="offset points",
                    ha="left", va="bottom", fontsize=FONT, color=INK, zorder=8)
    ax.set_xlabel("Quantity", fontsize=FONT, labelpad=8)
    ax.set_ylabel("Price", fontsize=FONT, labelpad=8)
    save(fig, out)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.parse_args()

    draw_class_demand(IMG / "class-demand.svg")
    draw_class_linear(IMG / "class-linear.svg")
    draw_class_elasticity(IMG / "class-elasticity.svg")
    draw_flat_steep(IMG / "flat-steep.svg")
    draw_q76(IMG / "core-q7-6.svg")
    draw_q78(IMG / "core-q7-8.svg")

    # Numbers quoted on the slides and in the worksheet, printed so they can
    # be checked against what is written in both.
    print("\nvalues used in slide and worksheet text")
    wtp = class_wtp()
    print(f"  class poll: {len(wtp)} answers, from ${wtp[-1]:.0f} to ${wtp[0]:.0f};"
          + "".join(f" {buyers(wtp, p)} at ${p}," for p in CLASS_PRICES))
    print(f"  straight line Q = {LIN_A} - {LIN_B:g} P; each ${LIN_STEP} rise:")
    for p0 in (LIN_EXAMPLE,) + LIN_ACTIVITY:
        p1 = p0 + LIN_STEP
        q0, q1 = class_linear(p0), class_linear(p1)
        dp, dq = 100 * (p1 - p0) / p0, 100 * (q1 - q0) / q0
        print(f"    ${p0} -> ${p1}: Q {q0:g} -> {q1:g}, %P = {dp:g}%, %Q = {dq:g}%,"
              f" eps = {-dq / dp:g}")
    print("  -(P/Q)(dQ/dP):")
    for p0 in POINT_ACTIVITY:
        q0 = class_linear(p0)
        print(f"    P = ${p0}: Q = {q0:g}, eps = {LIN_B * p0 / q0:g}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
